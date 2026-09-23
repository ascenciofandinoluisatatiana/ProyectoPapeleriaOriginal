"""
Generador de reportes en Excel usando openpyxl.
"""
import os
from datetime import datetime
from django.conf import settings
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter


def _carpeta_reportes():
    """Crea (si no existe) y retorna la carpeta donde se guardan los Excel."""
    carpeta = os.path.join(settings.MEDIA_ROOT, 'reportes')
    os.makedirs(carpeta, exist_ok=True)
    return carpeta


def _autoajustar_columnas(hoja):
    """Ajusta el ancho de cada columna según su contenido más largo."""
    for columna in hoja.columns:
        longitud_maxima = 0
        letra_columna = get_column_letter(columna[0].column)
        for celda in columna:
            try:
                if celda.value:
                    longitud_maxima = max(longitud_maxima, len(str(celda.value)))
            except Exception:
                pass
        hoja.column_dimensions[letra_columna].width = min(longitud_maxima + 3, 40)


def exportar_activos_excel(activos):
    """
    Genera un Excel con el listado de activos.
    Recibe un queryset de Activo y retorna la ruta del archivo generado.
    """
    wb = Workbook()
    hoja = wb.active
    hoja.title = 'Activos'

    encabezados = [
        'Código', 'Nombre', 'Categoría', 'Ubicación', 'Estado',
        'Marca', 'Modelo', 'Número de Serie', 'Responsable',
        'Fecha Adquisición', 'Valor Adquisición',
    ]
    hoja.append(encabezados)

    # Estilo del encabezado
    relleno = PatternFill(start_color='00324D', end_color='00324D', fill_type='solid')
    fuente = Font(color='FFFFFF', bold=True)
    for celda in hoja[1]:
        celda.fill = relleno
        celda.font = fuente
        celda.alignment = Alignment(horizontal='center')

    for activo in activos:
        hoja.append([
            activo.codigo_inventario,
            activo.nombre,
            str(activo.categoria) if activo.categoria else '',
            str(activo.ubicacion) if activo.ubicacion else '',
            str(activo.estado) if activo.estado else '',
            activo.marca or '',
            activo.modelo or '',
            activo.numero_serie or '',
            activo.responsable or '',
            activo.fecha_adquisicion.strftime('%d/%m/%Y') if activo.fecha_adquisicion else '',
            float(activo.valor_adquisicion) if activo.valor_adquisicion else 0,
        ])

    _autoajustar_columnas(hoja)
    hoja.freeze_panes = 'A2'

    nombre_archivo = f'activos_{datetime.now().strftime("%Y%m%d_%H%M%S")}.xlsx'
    ruta = os.path.join(_carpeta_reportes(), nombre_archivo)
    wb.save(ruta)

    return ruta


def exportar_logs_excel(logs):
    """
    Genera un Excel con el historial de auditoría.
    Recibe un queryset de Log y retorna la ruta del archivo generado.
    """
    wb = Workbook()
    hoja = wb.active
    hoja.title = 'Auditoría'

    encabezados = [
        'Fecha', 'Usuario', 'Acción', 'Modelo',
        'Objeto', 'Descripción', 'IP',
    ]
    hoja.append(encabezados)

    relleno = PatternFill(start_color='FF5722', end_color='FF5722', fill_type='solid')
    fuente = Font(color='FFFFFF', bold=True)
    for celda in hoja[1]:
        celda.fill = relleno
        celda.font = fuente
        celda.alignment = Alignment(horizontal='center')

    for log in logs:
        hoja.append([
            log.fecha_hora.strftime('%d/%m/%Y %H:%M:%S') if log.fecha_hora else '',
            log.usuario_nombre,
            log.get_accion_display(),
            log.modelo,
            log.objeto_repr or '',
            log.descripcion,
            log.ip_address or '',
        ])

    _autoajustar_columnas(hoja)
    hoja.freeze_panes = 'A2'

    nombre_archivo = f'logs_auditoria_{datetime.now().strftime("%Y%m%d_%H%M%S")}.xlsx'
    ruta = os.path.join(_carpeta_reportes(), nombre_archivo)
    wb.save(ruta)

    return ruta